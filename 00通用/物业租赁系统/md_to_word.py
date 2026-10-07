from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()

# Set page margins
sections = doc.sections
for section in sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = '宋体'
        run.font.size = Pt(11)
    return p

def add_heading2(doc, text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.name = '宋体'
        run.font.size = Pt(12)
    return p

def add_table(doc, rows_data, col_count):
    table = doc.add_table(rows=len(rows_data), cols=col_count)
    table.style = 'Table Grid'
    for i, row_data in enumerate(rows_data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    return table

# Title
p = doc.add_heading('家庭居室装修工程施工合同', 1)
for run in p.runs:
    run.font.name = '宋体'
doc.add_paragraph()

# Party info
add_para(doc, '发包方（甲方）：刘国云　　身份证号：42092319790724127X　　联系电话：17722592867')
add_para(doc, '承包方（乙方）：__________　　身份证号：__________　　联系电话：__________')
doc.add_paragraph('─' * 50)

# Article 1
add_heading2(doc, '第一条　工程概况')
add_para(doc, '1.1　工程地点：湖北省武汉市东西湖区惠景九龙湾小区9栋1单元15楼1501室')
add_para(doc, '1.2　房屋结构：框架')
add_para(doc, '1.3　装修面积：套内建筑面积约118平方米（最终以实测为准）')
add_para(doc, '1.4　工程内容：按本合同的约定进行家庭居室室内装修，包括设计、施工及材料采购')
add_para(doc, '1.5　承包方式：全包（包工包料）')
doc.add_paragraph('─' * 50)

# Article 2
add_heading2(doc, '第二条　工程工期')
add_para(doc, '2.1　工期总天数：________个日历天（不含设计时间）')
add_para(doc, '2.2　开工日期：2026年10月2日')
add_para(doc, '2.3　竣工日期：________年________月________日')
add_para(doc, '2.4　工期延误：因甲方原因、设计变更、不可抗力等因素导致工期顺延的，须经双方签字确认。因乙方原因每延误一日，按合同总价的________‰向甲方支付违约金。')
add_para(doc, '2.5　阶段工期：')
add_table(doc, [
    ['阶段', '工序', '工期（天）'],
    ['前期', '拆改、墙体新建', ''],
    ['水电', '水电改造、防水', ''],
    ['泥木', '贴砖、吊顶', ''],
    ['油漆', '墙面处理、刷漆', ''],
    ['安装', '灯具、开关、五金、保洁', ''],
], 3)
doc.add_paragraph('─' * 50)

# Article 3
add_heading2(doc, '第三条　工程造价与付款方式')
add_para(doc, '3.1　合同总造价：人民币（大写）捌万元整（¥80,000元）')
add_para(doc, '3.2　此价格为固定总价，包含本合同约定的全部施工项目和材料，除非双方书面签字确认变更，不得调整。')
add_para(doc, '3.3　付款方式（建议比例）：')
add_table(doc, [
    ['阶段', '付款比例', '金额（元）', '付款时间'],
    ['首付款', '30%', '24,000', '合同签订当日'],
    ['水电验收', '20%', '16,000', '水电隐蔽工程验收合格'],
    ['泥木验收', '30%', '24,000', '泥木工程验收合格'],
    ['竣工验收', '15%', '12,000', '整体工程验收合格'],
    ['尾款', '5%', '4,000', '入住30天后无质量问题'],
], 4)
add_para(doc, '3.4　付款账户须为乙方对公账户或乙方本人账户，收款后须提供有效票据。')
doc.add_paragraph('─' * 50)

# Article 4
add_heading2(doc, '第四条　设计方案与施工图纸')
add_para(doc, '4.1　设计方案由甲乙双方协商确定，双方共同参与方案讨论与确认。')
add_para(doc, '4.2　施工前须经双方确认最终方案，未经双方书面同意不得擅自施工或变更方案。')
doc.add_paragraph('─' * 50)

# Article 5
add_heading2(doc, '第五条　材料供应')
add_para(doc, '5.1　全包项目主材清单（详见附件《材料清单明细表》）：')
add_table(doc, [
    ['类别', '材料名称', '品牌/规格', '约定等级', '备注'],
    ['瓷砖', '地砖/墙砖', '', '', ''],
    ['地板', '实木/复合/石晶', '', '', ''],
    ['橱柜', '', '', '', ''],
    ['板材', '柜体板/背板/见光板', '', '', ''],
    ['柜体配件', '铰链/导轨/拉手/支撑等', '', '', ''],
    ['门', '卧室门/厨卫门/入户门', '', '', ''],
    ['卫浴', '马桶/花洒/浴室柜/龙头', '', '', ''],
    ['地漏', '厕所地漏', '', '', ''],
    ['三角阀', '冷热三角阀（洗衣机、洗手盆、热水器等）', '', '', ''],
    ['开关插座', '', '', '', ''],
    ['油漆/乳胶漆', '', '', '', ''],
    ['灯具', '', '', '', ''],
], 5)
doc.add_paragraph()
add_para(doc, '5.2　板材选定权：板材的品牌、型号、颜色、花色由甲方亲自选定，乙方不得擅自指定或替换。甲方须在橱柜深化设计阶段完成板材选定，并在《板材选定确认单》上签字确认，作为施工依据。')
add_para(doc, '5.3　材料验收：')
add_para(doc, '　• 所有材料进场须提前24小时通知甲方验收')
add_para(doc, '　• 甲方有权要求乙方提供材料合格证、检测报告')
add_para(doc, '　• 甲方验收签字后方可施工；未经签字擅自施工的，甲方有权要求返工')
add_para(doc, '5.4　材料品牌更换：乙方不得擅自更换约定品牌，确需更换须提前书面通知甲方并经同意，同档次同价位替换须有甲方书面确认。')
add_para(doc, '5.5　甲方未及时选定板材：若甲方未在约定时间内完成板材选定，导致工期延误，工期顺延责任由甲方自行承担，乙方不承担违约责任。')
doc.add_paragraph('─' * 50)

# Article 6
add_heading2(doc, '第六条　施工工艺与质量标准')
add_para(doc, '6.1　施工工艺须符合现行国家标准和地方规范，包括但不限于：')
add_para(doc, '　• 《住宅装饰装修工程施工规范》GB 50327')
add_para(doc, '　• 《建筑地面工程施工质量验收规范》GB 50209')
add_para(doc, '　• 《建筑内部装修设计防火规范》GB 50222')
add_para(doc, '6.2　主要工序质量要求：')
add_table(doc, [
    ['工序', '质量要求'],
    ['水电改造', '水管打压测试0.8MPa保持30分钟不掉压；电路阻值≥5MΩ；线管分色'],
    ['防水工程', '卫生间防水高度≥1.8m，厨房≥0.3m；闭水试验48小时无渗漏'],
    ['瓷砖铺贴', '空鼓率≤5%；缝隙均匀；阴阳角方正±2mm'],
    ['墙面处理', '2米靠尺检查误差≤3mm；无明显刷痕、流坠'],
    ['油漆工程', '无明显色差；无起皮、脱落、开裂'],
], 2)
add_para(doc, '6.3　隐蔽工程须在封闭前48小时通知甲方验收，未验收不得封闭。')
doc.add_paragraph('─' * 50)

# Article 7
add_heading2(doc, '第七条　工程变更')
add_para(doc, '7.1　工程变更须经双方签字确认《工程变更单》，明确变更内容、费用及工期影响。')
add_para(doc, '7.2　增项处理：')
add_para(doc, '　• 单次增项金额在________元以内的，乙方可先施工后确认')
add_para(doc, '　• 单次增项金额超过________元的，须提前书面确认后方可施工')
add_para(doc, '　• 增项累计金额不得超过合同总价的15%，超过部分须另签补充协议')
add_para(doc, '7.3　减项处理：按原报价扣除相应费用，并在结算时返还甲方。')
doc.add_paragraph('─' * 50)

# Article 8
add_heading2(doc, '第八条　工程验收')
add_para(doc, '8.1　分项验收：每个阶段工序完成后，甲方须在收到乙方验收通知后48小时内到场验收，逾期视为合格。')
add_para(doc, '8.2　竣工验收：整体工程完工后，乙方通知甲方验收，甲方须在7个工作日内组织验收，逾期视为合格。')
add_para(doc, '8.3　验收标准：以本合同约定、国家标准、地方规范及双方签字确认的图纸为准。')
add_para(doc, '8.4　验收记录：双方签字确认《分项验收记录表》《竣工验收单》，作为结算和保修依据。')
add_para(doc, '8.5　空气质量检测：竣工后须进行室内空气质量检测，甲醛等有害物质须符合GB 50325标准。')
doc.add_paragraph('─' * 50)

# Article 9
add_heading2(doc, '第九条　保修条款')
add_para(doc, '9.1　整体工程保修期：________年（自竣工验收合格之日起计算）')
add_para(doc, '9.2　分项保修期限：')
add_table(doc, [
    ['项目', '保修期限'],
    ['水电隐蔽工程', '10年'],
    ['防水工程/渗漏水', '10年'],
    ['墙面/墙皮脱落、开裂', '5年'],
    ['地面', '5年'],
    ['橱柜', '5年'],
    ['门窗、五金', '3年'],
    ['灯具、开关', '2年'],
    ['电器设备', '按厂家保修卡'],
], 2)
doc.add_paragraph()
add_para(doc, '9.3　重点质量问题保修：装修完工后，如出现以下质量问题，乙方须负责无偿维修或返工，造成甲方损失的由乙方赔偿：')
add_para(doc, '　• 渗漏水：包括但不限于防水层破损导致的墙面渗水、顶面渗水、邻里渗漏等')
add_para(doc, '　• 墙皮脱落：墙面乳胶漆、腻子层空鼓、脱落、起皮等')
add_para(doc, '　• 开裂：墙面、顶面裂纹、石膏板接缝开裂等')
add_para(doc, '　• 瓷砖空鼓、脱落：铺贴不牢导致的瓷砖空鼓、脱落')
add_para(doc, '　• 门窗、五金故障：因安装或材质问题导致的门窗开合不畅、锁具故障、五金损坏等')
add_para(doc, '9.4　保修范围：正常使用情况下出现的质量问题（非人为损坏、非不可抗力、非甲方擅自改动）。')
add_para(doc, '9.5　保修响应：乙方须在接到甲方保修通知后24小时内响应，72小时内上门维修。')
add_para(doc, '9.6　终身维护：保修期满后，乙方仍提供有偿维修服务，材料费按成本价收取。')
doc.add_paragraph('─' * 50)

# Article 10
add_heading2(doc, '第十条　双方权利与义务')
add_para(doc, '10.1　甲方权利与义务')
add_para(doc, '　• 及时支付合同约定的款项')
add_para(doc, '　• 提供装修期间的水电（费用自理）')
add_para(doc, '　• 提供施工所需的基础条件')
add_para(doc, '　• 有权对施工质量、材料进行监督检查')
add_para(doc, '　• 及时对乙方提出的验收、变更等事项作出答复')
add_para(doc, '10.2　乙方权利与义务')
add_para(doc, '　• 按合同约定和国家规范施工')
add_para(doc, '　• 对施工安全负责，承担施工现场的安全责任')
add_para(doc, '　• 负责协调各工种配合')
add_para(doc, '　• 按时完成各阶段工期')
add_para(doc, '　• 配合甲方进行验收')
add_para(doc, '　• 竣工后提供完整竣工图纸和保修凭证')
doc.add_paragraph('─' * 50)

# Article 11
add_heading2(doc, '第十一条　违约责任')
add_para(doc, '11.1　乙方擅自更换材料：须更换回约定品牌，并支付合同总价5%的违约金。')
add_para(doc, '11.2　工程质量不合格：乙方须在限期内返工至合格，返工费用由乙方承担；因此造成甲方损失的由乙方赔偿。')
add_para(doc, '11.3　乙方擅自解除合同：须退还甲方已付全部款项，并支付合同总价20%的违约金。')
doc.add_paragraph('─' * 50)

# Article 12
add_heading2(doc, '第十二条　争议解决')
add_para(doc, '12.1　本合同履行中发生的争议，双方应协商解决。')
add_para(doc, '12.2　协商不成的，任何一方可向合同签订地有管辖权的人民法院提起诉讼。')
add_para(doc, '12.3　诉讼期间，除争议事项外，双方应继续履行本合同其他条款。')
doc.add_paragraph('─' * 50)

# Article 13
add_heading2(doc, '第十三条　其他约定')
add_para(doc, '13.1　本合同一式两份，甲乙双方各执一份，具有同等法律效力。')
add_para(doc, '13.2　本合同自双方签字（盖章）之日起生效。')
add_para(doc, '13.3　附件为本合同不可分割的组成部分，与本合同具有同等法律效力：')
add_para(doc, '　• 附件1：《工程预算清单》')
add_para(doc, '　• 附件2：《材料清单明细表》')
add_para(doc, '　• 附件3：《施工图纸》')
add_para(doc, '　• 附件4：《工程变更单》（按需使用）')
add_para(doc, '　• 附件5：《分项验收记录表》')
add_para(doc, '　• 附件6：《甲供材料清单》（如有）')
doc.add_paragraph('─' * 50)

# Article 14
add_heading2(doc, '第十四条　补充条款')
add_para(doc, '（双方自行约定的其他事项）')
doc.add_paragraph()
add_para(doc, '__________')
add_para(doc, '__________')
add_para(doc, '__________')
doc.add_paragraph()

# Signature
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('甲方（签字）：__________　　　　　　　　乙方（签字）：__________')
run.font.name = '宋体'
run.font.size = Pt(11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('日期：________年________月________日')
run.font.name = '宋体'
run.font.size = Pt(11)

doc.add_paragraph('─' * 50)

# Usage Instructions
add_heading2(doc, '使用说明')
add_para(doc, '需要重点确认的事项：')
add_para(doc, '　1. 材料清单：逐项填写品牌、型号，避免"同档次"等模糊表述')
add_para(doc, '　2. 付款节点：建议验收后再付款，保留5%尾款')
add_para(doc, '　3. 违约金额：违约金过高可能被法院酌减，建议不超过合同总价的2‰/日')
add_para(doc, '　4. 增项上限：务必约定增项比例，防止施工中无限增项')
add_para(doc, '　5. 保修年限：建议整体工程≥2年，水电防水≥10年')
doc.add_paragraph()
add_para(doc, '常见陷阱提示：')
add_para(doc, '　⚠️ 口头承诺不写入合同')
add_para(doc, '　⚠️ 模糊报价（如"多退少补"无上限）')
add_para(doc, '　⚠️ 材料"以实际为准"无具体品牌型号')
add_para(doc, '　⚠️ 验收标准不明确')
add_para(doc, '　⚠️ 保修期从开工而非竣工算起')

# Save
output_path = 'D:/mywork/techdoc/00通用/物业租赁系统/家庭居室装修全包合同范本_v3.docx'
doc.save(output_path)
print(f'Word document saved: {output_path}')
